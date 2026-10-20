#!/usr/bin/env python3
"""Generate project documentation DOCX (design, prototypes, manual tests, unit tests)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "STT_Documentacia_Dizajn_i_Testing.docx"


def _style_doc(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    for i in range(1, 4):
        h = doc.styles[f"Heading {i}"]
        h.font.name = "Times New Roman"


def _h(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _p(doc: Document, text: str, bold: bool = False) -> None:
    run = doc.add_paragraph().add_run(text)
    run.bold = bold


def _bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = val


def section_design_requirements(doc: Document) -> None:
    _h(doc, "1. Требования к дизайну интерфейса приложения", 1)

    _h(doc, "1.1. Назначение и контекст продукта", 2)
    _p(
        doc,
        "STT App — локальное веб-приложение для macOS, предназначенное для одноразового "
        "сценария: пользователь загружает аудио- или видеофайл, получает расшифровку речи "
        "с разделением по спикерам (диаризация), редактирует текст, экспортирует результат "
        "в TXT и покидает сервис. История загрузок на главной странице не отображается. "
        "Администратор системы (локальный оператор) может через панель настроек изменять "
        "глобальные параметры: формальный tier модели и реальный лимит длительности файла.",
    )

    _h(doc, "1.2. Целевая аудитория и сценарии", 2)
    _bullets(
        doc,
        [
            "Конечный пользователь: сотрудник, исследователь, журналист — без технической подготовки.",
            "Администратор: владелец установки на Mac — настраивает лимиты для всех сессий.",
            "Основной сценарий (happy path): главная → загрузка → ожидание → результат → экспорт → «Новый файл».",
            "Дополнительный сценарий: администратор → настройки → логин → изменение лимита минут.",
        ],
    )

    _h(doc, "1.3. Принципы проектирования UX/UI", 2)
    _bullets(
        doc,
        [
            "Минимализм: на главной только зона загрузки; без лишней навигации и истории.",
            "Предсказуемость: единая тёмная тема, повторяющиеся паттерны карточек (card) и кнопок.",
            "Обратная связь: прогресс-бар с процентами и этапами; типизированные блоки ошибок.",
            "Локальность и доверие: подзаголовок «расшифровка с ролями · локально».",
            "Доступность: aria-label у иконки настроек; focus-visible у раскрывающихся блоков.",
            "Русскоязычный интерфейс; технические коды ошибок — вторичным мелким текстом.",
        ],
    )

    _h(doc, "1.4. Визуальный язык", 2)
    _p(doc, "Цветовая палитра (CSS-переменные реализации):", bold=True)
    _table(
        doc,
        ["Элемент", "Значение", "Применение"],
        [
            ["--bg", "#0f1218", "Фон страницы"],
            ["--surface", "#1a2030", "Карточки, модальные окна"],
            ["--border", "#2a3348", "Границы, dashed dropzone"],
            ["--text", "#e8ecf4", "Основной текст"],
            ["--muted", "#9aa7bd", "Подсказки, вторичный текст"],
            ["--accent", "#5b9dff", "Ссылки, акцентные кнопки, hover иконки"],
            ["--danger / alert-error", "#ffb8c0 на #2a1520", "Блоки ошибок"],
        ],
    )
    doc.add_paragraph()
    _p(doc, "Типографика:", bold=True)
    _bullets(
        doc,
        [
            "Системный стек: SF Pro Text, system-ui, -apple-system, sans-serif.",
            "Заголовок бренда STT: 1.25 rem, font-weight 700.",
            "Заголовок job (имя файла): 1.25 rem, semibold.",
            "Основной текст и поля ввода: 12 pt (наследование), line-height 1.5.",
            "Мелкий текст (small, hints): ~0.8125 rem, цвет muted.",
        ],
    )
    _p(doc, "Геометрия и отступы:", bold=True)
    _bullets(
        doc,
        [
            "Радиус скругления карточек и dropzone: 12 px (--radius).",
            "Максимальная ширина контента (app-shell): 1100 px, центрирование.",
            "Вертикальные gap между секциями: 1.25 rem (gap-lg).",
            "Кнопки: padding 0.5 rem × 1 rem; primary — синий фон, тёмный текст.",
        ],
    )

    _h(doc, "1.5. Экраны и обязательные элементы", 2)

    _h(doc, "1.5.1. Главная страница (/)", 3)
    _bullets(
        doc,
        [
            "Шапка: логотип-ссылка STT, подпись «расшифровка с ролями · локально».",
            "Справа в шапке: кнопка настроек (SVG-шестерёнка ~1.85 rem в кнопке ~2.85 rem).",
            "Центральная dropzone (min-height ~280 px): dashed border, hover/active — accent border.",
            "Текст: «Перетащите файл сюда»; форматы и лимит минут (из public settings API).",
            "Кнопка primary «Выбрать файл»; input file скрыт, accept audio/video + расширения.",
            "Блок ErrorNotice под dropzone при ошибках загрузки.",
            "Запрещено: список «Недавние», боковое меню, footer с историей.",
        ],
    )

    _h(doc, "1.5.2. Страница обработки и результата (/jobs/:id)", 3)
    _bullets(
        doc,
        [
            "Ссылка-кнопка «← Новый файл» (удаляет job на сервере).",
            "Имя исходного файла как заголовок.",
            "При processing: карточка ProcessingProgress — этап, %, полоса, сообщение backend.",
            "При failed: ErrorNotice с типом и текстом (parseJobError).",
            "При done: toolbar экспорта (формат TXT, склейка реплик, rediarize, copy, save).",
            "details «Имена спикеров»: summary с chevron, hint «нажмите, чтобы развернуть».",
            "Поиск по тексту с счётчиком найденных блоков.",
            "Две колонки (≥900 px): плеер | список реплик; иначе одна колонка.",
            "Плеер: native audio controls; подсветка active segment; click → seek.",
            "FAB справа снизу: scroll to top / bottom.",
            "Примечание diarization: «Обнаружен один спикер» при одном кластере.",
        ],
    )

    _h(doc, "1.5.3. Модальное окно настроек администратора", 3)
    _bullets(
        doc,
        [
            "Backdrop на весь viewport; клик вне — закрытие.",
            "Заголовок «Настройки системы», кнопка ×.",
            "Состояние login: логин, пароль, «Войти»; подсказка про сессию браузера.",
            "Состояние settings: select модели (лёгкая/средняя/тяжёлая), select минут (5–60).",
            "Кнопка «Применить для всех»; при 401 — сброс сессии и возврат к login.",
            "Без пояснения о «формальности» модели в UI (требование заказчика).",
        ],
    )

    _h(doc, "1.6. Состояния компонентов", 2)
    _table(
        doc,
        ["Компонент", "Состояния", "Визуальное поведение"],
        [
            ["Dropzone", "default, drag-over, disabled (loading)", "Активная рамка accent при drag"],
            ["Primary button", "default, hover, disabled", "Hover — accent-hover"],
            ["Progress bar", "running, done", "Creep до ~90% на ASR; 100% на done"],
            ["Transcript item", "default, segment-active, search-hit", "Accent border / green tint"],
            ["ErrorNotice", "visible + dismiss", "Красноватый фон, тип жирным"],
        ],
    )

    _h(doc, "1.7. Адаптивность и платформа", 2)
    _p(
        doc,
        "Приложение ориентировано на desktop-браузер на macOS (Chrome, Safari, Firefox). "
        "Минимальная ширина для двухколоночного layout результата — 900 px. "
        "Мобильная вёрстка не является приоритетом MVP, но touch-target кнопок ≥44 px желателен для FAB.",
    )

    _h(doc, "1.8. Нефункциональные требования к UI", 2)
    _bullets(
        doc,
        [
            "Время до первого контента (FCP): зависит от Vite dev/build; production — static assets.",
            "Не блокировать UI на upload: индикатор «Загрузка…» на кнопке выбора файла.",
            "Автосохранение правок transcript debounce ~600 ms; индикатор «Сохранение…/Сохранено».",
            "Session admin: sessionStorage, не localStorage — сброс при закрытии вкладки опционален.",
        ],
    )


def section_wireframe(doc: Document) -> None:
    _h(doc, "2. Создание каркасного прототипа интерфейса", 1)

    _h(doc, "2.1. Цели wireframe-этапа", 2)
    _p(
        doc,
        "Каркасный прототип фиксирует информационную архитектуру, потоки пользователя и "
        "расположение элементов без финальной визуальной стилизации. Для STT App wireframe "
        "служил основой для согласования отказа от истории загрузок и добавления админ-настроек.",
    )

    _h(doc, "2.2. Инструменты и артеfacts", 2)
    _bullets(
        doc,
        [
            "Figma / FigJam, Balsamiq, или ASCII-схемы в документации — допустимы на этапе идеи.",
            "В репозитории: логическая структура React-страниц (HomePage, JobPage) как отражение wireframe.",
            "Маршрутизация: react-router — «/» и «/jobs/:jobId» (два экрана MVP).",
        ],
    )

    _h(doc, "2.3. Карта экранов (screen map)", 2)
    _p(doc, "[Главная /] --upload--> [Job /jobs/:id] --done--> редактирование", bold=True)
    _p(doc, "[Главная /] --⚙--> [Modal: Login] --ok--> [Modal: Settings] --save--> закрытие")
    _p(doc, "[Job] --Новый файл--> [Главная] (+ DELETE job)")

    _h(doc, "2.4. Wireframe: главная (текстовая схема)", 2)
    _p(
        doc,
        "+----------------------------------------------------------+\n"
        "| STT   расшифровка с ролями · локально            [⚙]    |\n"
        "+----------------------------------------------------------+\n"
        "|                                                          |\n"
        "|     +----------------------------------------------+     |\n"
        "|     |         Перетащите файл сюда                 |     |\n"
        "|     |   форматы ... · до N мин                     |     |\n"
        "|     |            [ Выбрать файл ]                  |     |\n"
        "|     +----------------------------------------------+     |\n"
        "|                                                          |\n"
        "+----------------------------------------------------------+",
    )

    _h(doc, "2.5. Wireframe: страница job (processing)", 2)
    _p(
        doc,
        "+----------------------------------------------------------+\n"
        "| ← Новый файл                                             |\n"
        "| filename.mp3                                             |\n"
        "| +------------------------------------------------------+ |\n"
        "| | Распознавание речи                          73%     | |\n"
        "| | [===========...............]                         | |\n"
        "| | сообщение backend                                    | |\n"
        "| +------------------------------------------------------+ |\n"
        "+----------------------------------------------------------+",
    )

    _h(doc, "2.6. Wireframe: страница job (result)", 2)
    _p(
        doc,
        "Toolbar: [экспорт format v] [склеить] [Пересчитать] [Копировать] [Save txt]\n"
        "[▼ Имена спикеров — нажмите, чтобы развернуть]\n"
        "[ Поиск: _____________ ]\n"
        "+-------------+----------------------------------------+\n"
        "| [ audio ]   | Спикер 1  0:12                         |\n"
        "|  controls   | [ textarea реплика ]                   |\n"
        "|             | ...                                    |\n"
        "+-------------+----------------------------------------+\n"
        "                                    [↑]  [↓]  (fixed BR)",
    )

    _h(doc, "2.7. User Flow (диаграмма словами)", 2)
    _bullets(
        doc,
        [
            "UF-01 Upload: выбор файла → POST /api/transcriptions → redirect /jobs/id.",
            "UF-02 Wait: SSE events → обновление stage/message → progress UI.",
            "UF-03 Edit: PATCH segments/speakers с debounce.",
            "UF-04 Export: POST export.txt → download blob.",
            "UF-05 Admin: login → PATCH admin/settings → public limit обновлён.",
        ],
    )

    _h(doc, "2.8. Критерии приёмки wireframe", 2)
    _bullets(
        doc,
        [
            "Согласованы все обязательные поля на каждом экране.",
            "Нет экрана «история» и «профиль пользователя».",
            "Поток администратора отделён от потока пользователя (модалка).",
            "Предусмотрены состояния ошибки на upload и processing.",
        ],
    )


def section_hifi_prototype(doc: Document) -> None:
    _h(doc, "3. Создание прототипа интерфейса (с дизайном)", 1)

    _h(doc, "3.1. Переход от wireframe к UI", 2)
    _p(
        doc,
        "Высокодетализированный прототип реализован как рабочий frontend на React 19 + Vite 6 + TypeScript. "
        "Стили — custom CSS (frontend/src/styles.css) в тёмной теме. Это interactive high-fidelity prototype, "
        "подключённый к реальному FastAPI backend.",
    )

    _h(doc, "3.2. Структура frontend", 2)
    _table(
        doc,
        ["Файл / компонент", "Назначение"],
        [
            ["App.tsx", "Shell, header, routing, settings modal trigger"],
            ["pages/HomePage.tsx", "Dropzone, upload, public max minutes"],
            ["pages/JobPage.tsx", "Progress, result, search, export, delete on exit"],
            ["SettingsModals.tsx", "Admin login + settings"],
            ["ProcessingProgress.tsx", "Animated progress bar"],
            ["ErrorNotice.tsx", "Typed errors"],
            ["ScrollJumpButtons.tsx", "FAB scroll"],
            ["AudioPlayerControls.tsx", "Native audio element"],
            ["api.ts", "REST + SSE + error parsing"],
        ],
    )

    _h(doc, "3.3. Визуальные решения high-fidelity", 2)
    _bullets(
        doc,
        [
            "Тёмная тема color-scheme: light dark — комфорт при длительном чтении расшифровки.",
            "Dropzone hero: крупная центральная область, dashed → solid accent при drag.",
            "Карточки transcript-item с nested textarea — WYSIWYG-редактирование реплик.",
            "Модалка настроек: overlay 55% black, card max-width 440 px.",
            "Иконка settings: SVG stroke cog (Heroicons-style), hover accent.",
        ],
    )

    _h(doc, "3.4. Интерактивность прототипа", 2)
    _bullets(
        doc,
        [
            "Drag-and-drop и file picker на главной.",
            "SSE-driven progress (subscribeJobEvents) с fix ложного network error на done.",
            "timeupdate на audio → подсветка segment-active.",
            "details/summary для спикеров с rotate chevron.",
            "sessionStorage для admin token.",
        ],
    )

    _h(doc, "3.5. Связь с backend API", 2)
    _p(doc, "Vite proxy /api → 127.0.0.1:8000. Основные endpoints описаны в docs/API.md.")

    _h(doc, "3.6. Сборка и демонстрация", 2)
    _bullets(
        doc,
        [
            "Dev: npm run dev → http://127.0.0.1:5173",
            "Production build: npm run build → frontend/dist",
            "Скриншоты для отчёта: главная, processing, result, settings modal",
        ],
    )


def section_manual_tests(doc: Document) -> None:
    _h(doc, "4. Разработка тестовых сценариев для ручного тестирования", 1)

    _h(doc, "4.1. Тестовое окружение", 2)
    _bullets(
        doc,
        [
            "macOS, Python 3.11+ venv в backend, Node 20+ для frontend.",
            "backend/.env: ADMIN_USERNAME/PASSWORD, опционально HF_TOKEN.",
            "Запуск: ./scripts/run_backend.sh и ./scripts/run_frontend.sh.",
            "Тестовые файлы: короткий WAV/MP3 (<1 мин), диалог 2 спикера (~5 мин), видео MKV с аудио, файл > лимита минут.",
        ],
    )

    _h(doc, "4.2. Матрица сценариев", 2)
    rows = [
        ["TC-01", "Главная", "Открыть /", "Dropzone, кнопка выбора, ⚙, нет истории", "P0"],
        ["TC-02", "Upload audio", "MP3 < лимита", "Redirect job, status processing", "P0"],
        ["TC-03", "Upload video", "MKV с аудио", "Успешная обработка", "P1"],
        ["TC-04", "No audio", "Видео без дорожки", "invalid_audio", "P1"],
        ["TC-05", "Too long", "Файл > max_duration", "audio_too_long + понятный текст", "P0"],
        ["TC-06", "Progress", "Длинный файл", "Полоса до ~90% на ASR, 100% на done", "P1"],
        ["TC-07", "SSE done", "Дождаться конца", "Нет «потеряно соединение»", "P1"],
        ["TC-08", "Result player", "Клик time tag", "Seek + play, подсветка", "P0"],
        ["TC-09", "Edit text", "Изменить реплику", "Автосохранение, Сохранено", "P0"],
        ["TC-10", "Speakers", "Развернуть details", "Переименование в export", "P1"],
        ["TC-11", "Search", "Ввести substring", "Фильтр + search-hit", "P1"],
        ["TC-12", "Scroll FAB", "↑ ↓", "Smooth scroll", "P2"],
        ["TC-13", "Export dialogue", "Save txt", "Формат Спикер: текст", "P0"],
        ["TC-14", "Export timestamps", "Save txt", "[HH:MM:SS] prefix", "P1"],
        ["TC-15", "Rediarize", "Кнопка", "Обновление speaker blocks", "P2"],
        ["TC-16", "New file", "← Новый файл", "Главная, job удалён", "P1"],
        ["TC-17", "Admin login fail", "Неверный пароль", "admin_invalid_credentials", "P1"],
        ["TC-18", "Admin login ok", "admin/admin", "Форма настроек", "P0"],
        ["TC-19", "Change limit", "15 min + save", "Главная показывает «до 15 мин»", "P0"],
        ["TC-20", "Admin session", "Restart backend, open ⚙", "Форма login, не error+settings", "P1"],
        ["TC-21", "Model tier", "Выбрать heavy, save", "Сохраняется в system_settings.json", "P2"],
    ]
    _table(doc, ["ID", "Область", "Действие", "Ожидание", "Приоритет"], rows)

    _h(doc, "4.3. Детальные шаги (примеры P0)", 2)

    _h(doc, "TC-02: Загрузка аудио", 3)
    _bullets(
        doc,
        [
            "Предусловие: backend и frontend запущены, лимит ≥ длительности файла.",
            "Шаг 1: Открыть главную.",
            "Шаг 2: Нажать «Выбрать файл», выбрать sample.mp3.",
            "Шаг 3: Убедиться в переходе на /jobs/{uuid}.",
            "Шаг 4: Дождаться status=done.",
            "Ожидание: список сегментов не пуст; нет необработанной 500.",
        ],
    )

    _h(doc, "TC-18–19: Админ настройки", 3)
    _bullets(
        doc,
        [
            "Шаг 1: На главной нажать ⚙.",
            "Шаг 2: Ввести admin / admin → Войти.",
            "Шаг 3: Установить max duration 15 мин → Применить.",
            "Шаг 4: Закрыть модалку, обновить главную.",
            "Ожидание: текст «до 15 мин»; upload 20 мин файла → audio_too_long.",
        ],
    )

    _h(doc, "4.4. Негативные и граничные случаи", 2)
    _bullets(
        doc,
        [
            "Backend остановлен → upload: network/unknown error.",
            "Очень большой файл → file_too_large (413).",
            "Пустой/битый файл → invalid_audio.",
            "Закрытие вкладки во время processing → job может остаться на сервере (известное ограничение).",
        ],
    )

    _h(doc, "4.5. Регрессионный чек-лист перед релизом", 2)
    _bullets(
        doc,
        [
            "npm run build без ошибок TypeScript.",
            "pytest backend/tests — все green.",
            "Ручной TC-01, TC-02, TC-08, TC-13, TC-18, TC-19.",
        ],
    )


def section_unit_tests(doc: Document) -> None:
    _h(doc, "5. Разработка модульных тестов", 1)

    _h(doc, "5.1. Стратегия модульного тестирования", 2)
    _p(
        doc,
        "Модульные тесты сфокусированы на backend (Python/pytest): детерминированная бизнес-логика "
        "без GPU и без тяжёлых ML-моделей. Интеграционные тесты Whisper/pyannote выполняются вручную "
        "или в отдельной CI-джобе с моделями. Frontend на MVP покрывается build (tsc) и ручными сценариями; "
        "Vitest/React Testing Library — рекомендуемое расширение.",
    )

    _h(doc, "5.2. Инфраструктура", 2)
    _bullets(
        doc,
        [
            "Runner: pytest ≥8.",
            "Запуск: cd backend && .venv/bin/pytest tests/ -v",
            "Каталог: backend/tests/",
            "Зависимости dev: httpx (для будущих API-тестов), pytest.",
        ],
    )

    _h(doc, "5.3. Существующие тест-модули", 2)

    _h(doc, "5.3.1. test_export.py", 3)
    _p(doc, "Модуль: app.services.export.export_txt")
    _bullets(
        doc,
        [
            "Проверяет формат dialogue: «Спикер: текст».",
            "Проверяет timestamps: префикс [00:00:01].",
            "Использует TranscriptResult с двумя сегментами и speakers.",
        ],
    )

    _h(doc, "5.3.2. test_system_settings.py", 3)
    _p(doc, "Модуль: app.services.system_settings")
    _bullets(
        doc,
        [
            "monkeypatch data_dir на tmp_path.",
            "save_system_settings + load + get_max_duration_sec.",
            "Проверка max_duration_minutes=15 → 900 сек.",
        ],
    )

    _h(doc, "5.4. Рекомендуемые дополнительные модульные тесты", 2)
    _table(
        doc,
        ["Модуль", "Тест-кейс", "Приоритет"],
        [
            ["audio.py", "probe_duration на temp WAV без ffprobe (mock)", "P1"],
            ["audio.py", "assert_has_audio_ffmpeg detects Audio stream line", "P1"],
            ["api_errors.py", "api_error structure in HTTPException.detail", "P2"],
            ["admin_session.py", "create_session / validate / expiry", "P1"],
            ["transcription._assign_speakers", "overlap + midpoint fallback", "P1"],
            ["diarization._estimate_num_speakers", "synthetic embeddings mock", "P2"],
            ["routers (TestClient)", "POST upload rejects too long (mock probe)", "P1"],
            ["routers (TestClient)", "admin login 401/200", "P0"],
        ],
    )

    _h(doc, "5.5. Пример кода теста (admin login)", 2)
    _p(
        doc,
        "from fastapi.testclient import TestClient\n"
        "from app.main import app\n\n"
        "client = TestClient(app)\n\n"
        "def test_admin_login_ok():\n"
        "    r = client.post('/api/admin/login', json={'username':'admin','password':'admin'})\n"
        "    assert r.status_code == 200\n"
        "    assert 'token' in r.json()",
    )

    _h(doc, "5.6. Изоляция и fixtures", 2)
    _bullets(
        doc,
        [
            "tmp_path для system_settings и файловой I/O.",
            "monkeypatch settings.data_dir — не трогать реальный data/.",
            "Не вызывать WhisperModel в unit-тестах по умолчанию.",
            "Fixture TestClient с override dependencies при необходимости.",
        ],
    )

    _h(doc, "5.7. Критерии покрытия MVP", 2)
    _bullets(
        doc,
        [
            "100% покрытие export_txt форматов (dialogue, timestamps, merge_adjacent).",
            "Покрытие save/load system settings и allowed duration validation.",
            "≥1 тест на structured API errors для upload validation.",
            "Документированная команда запуска в README и данном DOCX.",
        ],
    )

    _h(doc, "5.8. CI (рекомендация)", 2)
    _p(
        doc,
        "GitHub Actions: job backend-test на push — setup Python, pip install -r requirements.txt, "
        "pytest tests/. Job frontend-build — npm ci && npm run build. ML smoke — optional scheduled.",
    )


def main() -> None:
    doc = Document()
    _style_doc(doc)

    title = doc.add_heading("STT App — документация по дизайну, прототипам и тестированию", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "Проект: my_stt_app — локальная расшифровка аудио/видео с разделением по спикерам.\n"
        "Версия документа: 1.0. Стек: FastAPI, React, faster-whisper, Resemblyzer/pyannote."
    )
    doc.add_paragraph()

    section_design_requirements(doc)
    doc.add_page_break()
    section_wireframe(doc)
    doc.add_page_break()
    section_hifi_prototype(doc)
    doc.add_page_break()
    section_manual_tests(doc)
    doc.add_page_break()
    section_unit_tests(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Written: {OUTPUT}")


if __name__ == "__main__":
    main()
